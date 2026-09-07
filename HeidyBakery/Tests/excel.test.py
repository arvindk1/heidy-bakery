"""Independent recalculation of actual native XLSX output. Requires openpyxl and LibreOffice."""
import json, math, pathlib, subprocess, sys
from openpyxl import load_workbook
root=pathlib.Path(__file__).resolve().parents[1]
work=pathlib.Path(sys.argv[1]).resolve(); work.mkdir(parents=True,exist_ok=True)
subprocess.run(['node',str(root/'Tests/excel-fixtures.cjs'),str(work)],check=True)
exe=str(root/'Heidy Bakery.app/Contents/MacOS/HeidyBakery')
expected=json.loads((work/'expected.json').read_text())
for name in expected: subprocess.run([exe,'--export-fixture',str(work/(name+'.json')),str(work/(name+'.xlsx'))],check=True,stdout=subprocess.DEVNULL)
recalc=work/'recalculated';recalc.mkdir(exist_ok=True)
result=subprocess.run(['soffice','--headless','-env:UserInstallation='+ (work/'lo-profile').as_uri(),'--convert-to','xlsx','--outdir',str(recalc),*[str(work/(n+'.xlsx'))for n in expected]],capture_output=True,text=True)
if result.returncode: raise RuntimeError(result.stderr)
for name,recipes in expected.items():
    book=load_workbook(recalc/(name+'.xlsx'),data_only=True)
    errors=[f'{s.title}!{c.coordinate}: {c.value}'for s in book for row in s for c in row if c.data_type=='e']
    assert not errors,(name,errors[:10])
    for n,r in enumerate(recipes,2):
        for col,want in enumerate(r['values'],8):
            got=book['Price list'].cell(n,col).value
            if want is None: assert got in (None,''),(name,r['name'],col,got,want)
            else: assert isinstance(got,(float,int)) and math.isclose(got,want,rel_tol=1e-10,abs_tol=1e-10),(name,r['name'],col,got,want)
print(f'Independent Excel parity passed for {len(expected)} edited-workbook scenarios; no formula errors.')
